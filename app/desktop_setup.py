"""Resumable dependency preparation shared by Windows Setup and desktop Settings."""

import hashlib
import json
import os
import shutil
import subprocess
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from urllib.request import Request, urlopen

import psutil
from filelock import FileLock

from . import core, lux_bootstrap, media, processes
from .platform_paths import desktop_directories

OFFICE_VERSION = "26.8.1"
OFFICE_NAME = f"LibreOffice_{OFFICE_VERSION}_Win_x86-64.msi"
OFFICE_URL = (
    f"https://download.documentfoundation.org/libreoffice/stable/{OFFICE_VERSION}/win/x86_64/{OFFICE_NAME}"
)
OFFICE_SHA256 = "c6298b10cfa1748bcbedad94c62ce771e240e5021bfd6d7a967cd571b04b5312"


def status_path():
    return core.DATA / "setup-status.json"


def read_status():
    try:
        value = json.loads(status_path().read_text(encoding="utf-8"))
        if value.get("status") == "running":
            try:
                active = psutil.Process(value["pid"]).create_time() == value["process_started"]
            except (psutil.Error, KeyError):
                active = False
            if not active:
                value.update(
                    status="failed", message="Setup was interrupted. Retry to finish preparing components."
                )
        return value
    except (OSError, ValueError):
        return {"status": "not_started", "message": "Prepare local narration and PowerPoint support."}


def report(status, message):
    value = {
        "status": status,
        "message": message,
        "updated_at": time.time(),
        "pid": os.getpid(),
        "process_started": psutil.Process().create_time(),
    }
    path = status_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(value), encoding="utf-8")
    temporary.replace(path)
    print(message, flush=True)


def verified_download(url, destination, checksum):
    if destination.is_file():
        with destination.open("rb") as cached:
            if hashlib.file_digest(cached, "sha256").hexdigest() == checksum:
                return
    partial = destination.with_suffix(destination.suffix + ".part")
    try:
        with urlopen(Request(url, headers={"Range": "bytes=0-0"}), timeout=60) as probe:
            ranged = getattr(probe, "status", 200) == 206
            total = int(probe.headers.get("Content-Range", "0/0").split("/")[-1]) if ranged else 0
            source = probe.geturl() if ranged else url
        if ranged and total:
            download_ranges(source, partial, total, checksum)
        else:
            download_sequential(url, partial)
        with partial.open("rb") as download:
            if hashlib.file_digest(download, "sha256").hexdigest() != checksum:
                shutil.rmtree(
                    partial.with_name(partial.name + ".chunks-" + checksum[:12]), ignore_errors=True
                )
                raise RuntimeError("LibreOffice download checksum did not match. Please retry setup.")
        partial.replace(destination)
        shutil.rmtree(partial.with_name(partial.name + ".chunks-" + checksum[:12]), ignore_errors=True)
    finally:
        partial.unlink(missing_ok=True)


def download_sequential(url, partial):
    with urlopen(url, timeout=60) as response, partial.open("wb") as output:
        total = int(response.headers.get("Content-Length", 0))
        received = 0
        last_report = 0
        while chunk := response.read(1024 * 1024):
            output.write(chunk)
            received += len(chunk)
            if time.monotonic() - last_report > 5:
                progress = f"{received // 1048576} MB" + (f" of {total // 1048576} MB" if total else "")
                report("running", f"Downloading PowerPoint support: {progress}…")
                last_report = time.monotonic()


def download_ranges(url, partial, total, checksum, chunk_size=8 * 1024 * 1024):
    chunks = partial.with_name(partial.name + ".chunks-" + checksum[:12])
    chunks.mkdir(exist_ok=True)
    lock = threading.Lock()
    completed = 0

    def fetch(first):
        nonlocal completed
        last = min(first + chunk_size, total) - 1
        path = chunks / str(first)
        if not path.is_file() or path.stat().st_size != last - first + 1:
            temporary = path.with_suffix(".tmp")
            for attempt in range(3):
                try:
                    request = Request(url, headers={"Range": f"bytes={first}-{last}"})
                    with urlopen(request, timeout=60) as response, temporary.open("wb") as output:
                        expected_range = f"bytes {first}-{last}/{total}"
                        if response.status != 206 or response.headers.get("Content-Range") != expected_range:
                            raise RuntimeError("The download server returned an invalid range. Retry setup.")
                        shutil.copyfileobj(response, output, length=1024 * 1024)
                    if temporary.stat().st_size != last - first + 1:
                        raise RuntimeError("A download chunk was incomplete. Retry setup.")
                    temporary.replace(path)
                    break
                except (OSError, RuntimeError):
                    temporary.unlink(missing_ok=True)
                    if attempt == 2:
                        raise
                    time.sleep(2**attempt)
        with lock:
            completed += path.stat().st_size
            report(
                "running",
                f"Downloading PowerPoint support: {completed // 1048576} MB of {total // 1048576} MB…",
            )
        return path

    with ThreadPoolExecutor(max_workers=8) as pool:
        paths = list(pool.map(fetch, range(0, total, chunk_size)))
    with partial.open("wb") as output:
        for path in paths:
            with path.open("rb") as piece:
                shutil.copyfileobj(piece, output)


def run_child(command, **options):
    with subprocess.Popen(
        command,
        stdin=subprocess.DEVNULL,
        stdout=sys.stdout,
        stderr=sys.stderr,
        **processes.background_options(),
        **options,
    ) as process:
        try:
            result = process.wait(timeout=1800)
        except BaseException:
            processes.stop_process_tree(process)
            raise
    if result:
        raise RuntimeError(f"Dependency preparation exited with code {result}. See data/logs/desktop.log.")


def prepare_office():
    if media.soffice():
        return
    if sys.platform != "win32":
        return  # The Windows installer owns automatic LibreOffice provisioning.
    cache = desktop_directories()["cache"]
    downloads = cache / "downloads"
    downloads.mkdir(parents=True, exist_ok=True)
    installer = downloads / OFFICE_NAME
    verified_download(OFFICE_URL, installer, OFFICE_SHA256)
    target = cache / "libreoffice" / OFFICE_VERSION
    target.mkdir(parents=True, exist_ok=True)
    log = core.DATA / "logs" / "libreoffice-setup.log"
    log.parent.mkdir(parents=True, exist_ok=True)
    report("running", "Preparing private PowerPoint renderer…")
    # An administrative image extracts the renderer into our writable cache.
    # Do not change the user's Office defaults or uninstall their office suite.
    run_child(
        ["msiexec.exe", "/a", str(installer), "/qn", "/norestart", f"TARGETDIR={target}", "/L*v", str(log)]
    )
    candidates = list(target.rglob("program/soffice.exe"))
    if not candidates:
        raise RuntimeError("LibreOffice extraction did not produce its renderer. Retry setup.")
    renderer = candidates[0]
    from .desktop_settings import DesktopSettings, save_settings

    # Persist through the same atomic writer used for desktop configuration.
    save_settings(DesktopSettings(), extra_updates={"SOFFICE_PATH": str(renderer)})
    run_child([str(renderer), "--headless", "--version"])
    installer.unlink(missing_ok=True)


def prepare_all():
    lock = FileLock(core.DATA / "setup.lock")
    with lock.acquire(timeout=0):
        try:
            report("running", "Preparing Python and LuxTTS dependencies. Internet is required…")
            python, worker = lux_bootstrap.prepare_runtime()
            report("running", "Downloading and checking narration models. This can take several minutes…")
            environment = os.environ.copy()
            environment["PYTHONPATH"] = str(worker)
            run_child(
                [str(python), "-c", "from app.lux_server import LuxEngine; LuxEngine()"],
                cwd=worker,
                env=environment,
            )
            report("running", "Preparing PowerPoint support…")
            prepare_office()
            message = (
                "Local narration and PowerPoint support are ready."
                if media.soffice()
                else ("Local narration is ready. Install LibreOffice for PowerPoint, or upload PDFs.")
            )
            report("complete", message)
        except Exception:
            report(
                "failed",
                "Setup did not finish. Check your internet connection and retry. Details: data/logs/desktop.log.",
            )
            raise


def verify_offline():
    """Exercise the downloaded models and a real PPTX without network access."""
    import tempfile

    import fitz
    from pptx import Presentation

    python, worker = lux_bootstrap.prepare_runtime()
    environment = os.environ.copy()
    environment.update(PYTHONPATH=str(worker), HF_HUB_OFFLINE="1", TRANSFORMERS_OFFLINE="1")
    run_child(
        [str(python), "-c", "from app.lux_server import LuxEngine; LuxEngine()"], cwd=worker, env=environment
    )
    with tempfile.TemporaryDirectory(prefix="setup-verification-", dir=core.DATA) as directory:
        from pathlib import Path

        folder = Path(directory)
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "SlideVoice setup check"
        source = folder / "setup-check.pptx"
        presentation.save(source)
        pdf = media.pptx_to_pdf(source, folder)
        with fitz.open(pdf) as document:
            if document.page_count != 1:
                raise RuntimeError("PowerPoint renderer verification failed.")
    print("Offline narration model loading and PowerPoint conversion passed.", flush=True)
