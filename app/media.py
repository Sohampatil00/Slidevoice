import math
import os
import shutil
import subprocess
import sys
import tempfile
import wave
import zipfile
from array import array
from pathlib import Path

import imageio_ffmpeg
import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation

from .core import MAX_SLIDES, WorkflowError
from .processes import background_options

FPS = 30


def ffmpeg():
    return os.getenv("FFMPEG_PATH") or shutil.which("ffmpeg") or imageio_ffmpeg.get_ffmpeg_exe()


def soffice():
    configured = os.getenv("SOFFICE_PATH")
    candidates = [
        configured,
        shutil.which("soffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        str(Path.home() / "Applications/LibreOffice.app/Contents/MacOS/soffice"),
        *[
            str(Path(os.getenv(key, default)) / "LibreOffice/program/soffice.exe")
            for key, default in (
                ("ProgramFiles", "C:/Program Files"),
                ("ProgramFiles(x86)", "C:/Program Files (x86)"),
            )
        ],
    ]
    return next((path for path in candidates if path and Path(path).is_file()), None)


def run_command(arguments, timeout=300):
    try:
        completed = subprocess.run(
            arguments, capture_output=True, timeout=timeout, check=False, **background_options()
        )
    except subprocess.TimeoutExpired as error:
        raise WorkflowError("Processing timed out. Try a smaller presentation.") from error
    except OSError as error:
        raise WorkflowError("A required media tool could not be started. Check Settings.") from error
    if completed.returncode:
        raise WorkflowError("Media processing failed. Check the file is valid and not password-protected.")
    return completed


def validate_pptx(source):
    try:
        with zipfile.ZipFile(source) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(entry.file_size for entry in entries) > 300 * 1024 * 1024:
                raise WorkflowError("This presentation expands beyond the 300 MB processing limit.")
            if "ppt/presentation.xml" not in archive.namelist():
                raise WorkflowError("This is not a valid PPTX presentation.")
    except zipfile.BadZipFile as error:
        raise WorkflowError("This is not a valid PPTX presentation.") from error


def pptx_to_pdf(source, folder):
    executable = soffice()
    if not executable:
        raise WorkflowError(
            "PPTX rendering needs LibreOffice. Install it, or export your PPT as PDF and upload that."
        )
    validate_pptx(source)
    with tempfile.TemporaryDirectory(prefix="slidevoice-office-") as profile:
        run_command(
            [
                executable,
                f"-env:UserInstallation={Path(profile).as_uri()}",
                "--headless",
                "--convert-to",
                "pdf:impress_pdf_Export",
                "--outdir",
                str(folder),
                str(source),
            ]
        )
    target = folder / (source.stem + ".pdf")
    if not target.is_file():
        raise WorkflowError("PowerPoint conversion failed. Export to PDF in PowerPoint and try that file.")
    return target


def slide_record(position, image, title="", text="", notes=""):
    return {
        "id": f"slide-{position:03}",
        "number": position,
        "image": image,
        "title": title.strip()[:160] or f"Slide {position}",
        "text": text[:18000],
        "notes": notes[:12000],
        "script": "",
        "audio": None,
        "duration": None,
        "audio_signature": None,
        "audio_version": None,
        "audio_source": None,
    }


def import_document(source, folder):
    metadata = []
    document_path = source
    if source.suffix.lower() == ".pptx":
        validate_pptx(source)
        presentation = Presentation(source)
        if len(presentation.slides) > MAX_SLIDES:
            raise WorkflowError(f"Please limit presentations to {MAX_SLIDES} slides.")
        for slide in presentation.slides:
            text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
            notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
            title = slide.shapes.title.text if slide.shapes.title is not None else ""
            metadata.append((title, text, notes))
        document_path = pptx_to_pdf(source, folder)
    with fitz.open(document_path) as document:
        if document.needs_pass:
            raise WorkflowError("Password-protected PDFs are not supported.")
        if not 0 < len(document) <= MAX_SLIDES:
            raise WorkflowError(f"Please upload between 1 and {MAX_SLIDES} slides.")
        if metadata and len(metadata) != len(document):
            raise WorkflowError(
                "PPTX rendering changed the slide count. Export all slides to PDF in PowerPoint, then upload the PDF."
            )
        slides = []
        for position, page in enumerate(document, 1):
            if page.rect.width <= 0 or page.rect.height <= 0:
                raise WorkflowError("A slide has invalid dimensions.")
            scale = min(1920 / page.rect.width, 1080 / page.rect.height)
            image_name = f"slide-{position:03}.png"
            page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).save(folder / image_name)
            text = page.get_text()
            title = next((line for line in text.splitlines() if line.strip()), "")
            details = metadata[position - 1] if metadata else (title, text, "")
            slides.append(slide_record(position, image_name, *details))
        return slides


def import_image(source, folder, position, title):
    filename = f"slide-{position:03}.png"
    with Image.open(source) as original:
        original.load()
        image = ImageOps.exif_transpose(original).convert("RGB")
        image.thumbnail((1920, 1080))
        image.save(folder / filename)
    return slide_record(position, filename, title=title)


class SilentAudioError(WorkflowError):
    pass


def validate_narration_audio(path):
    """Require some audible signal in our normalized mono PCM recordings.

    Count 20 ms windows above -60 dBFS, ignoring DC offsets. Requiring
    100 ms of signal rejects silence and isolated clicks while allowing pauses.
    This checks audibility, not whether the model spoke the correct words.
    """
    try:
        with wave.open(str(path), "rb") as recording:
            if (recording.getnchannels(), recording.getsampwidth(), recording.getframerate()) != (
                1,
                2,
                48000,
            ):
                raise WorkflowError("Narration audio is invalid. Generate or upload it again.")
            duration = recording.getnframes() / recording.getframerate()
            if not 0.1 <= duration <= 1800:
                raise WorkflowError("Audio must be between 0.1 seconds and 30 minutes per slide.")
            audible_samples = 0
            while frames := recording.readframes(960):
                samples = array("h", frames)
                if sys.byteorder != "little":
                    samples.byteswap()
                mean = sum(samples) / len(samples)
                variance = sum(sample * sample for sample in samples) / len(samples) - mean * mean
                if variance >= (32768 * 0.001) ** 2:
                    audible_samples += len(samples)
                if audible_samples >= 4800:
                    return duration
    except (OSError, EOFError, wave.Error, ValueError) as error:
        raise WorkflowError("Narration audio is missing or invalid. Generate or upload it again.") from error
    raise SilentAudioError(
        "Narration audio is silent or too quiet. Regenerate this slide or upload an audible recording."
    )


def normalize_audio(source, destination, speed=1.0):
    temporary = destination.with_name(destination.stem + "-normalizing.wav")
    try:
        run_command(
            [
                ffmpeg(),
                "-nostdin",
                "-y",
                "-v",
                "error",
                "-i",
                str(source),
                "-vn",
                *(["-af", f"atempo={speed}"] if speed != 1.0 else []),
                "-ac",
                "1",
                "-ar",
                "48000",
                "-c:a",
                "pcm_s16le",
                str(temporary),
            ]
        )
        duration = validate_narration_audio(temporary)
        temporary.replace(destination)
        return duration
    finally:
        temporary.unlink(missing_ok=True)


def timeline_for(slides, pause):
    if not math.isfinite(pause) or not 0 <= pause <= 5:
        raise WorkflowError("Pause must be between 0 and 5 seconds.")
    cursor_frames = 0
    timeline = []
    for slide in slides:
        duration = slide.get("duration")
        if not slide.get("audio") or not duration or not math.isfinite(duration):
            raise WorkflowError(f"Generate or upload narration for slide {slide['number']} first.")
        frames = math.ceil((duration + pause) * FPS)
        timeline.append(
            {
                "slide": slide["number"],
                "title": slide["title"],
                "start": cursor_frames / FPS,
                "end": (cursor_frames + frames) / FPS,
                "audio_duration": duration,
                "frames": frames,
                "script": slide["script"],
            }
        )
        cursor_frames += frames
    return timeline


def render_video(folder, slides, pause, resolution, progress):
    dimensions = {"1080p": (1920, 1080), "720p": (1280, 720)}
    width, height = dimensions[resolution]
    timeline = timeline_for(slides, pause)
    for slide in slides:
        try:
            validate_narration_audio(folder / slide["audio"])
        except WorkflowError as error:
            raise WorkflowError(f"Slide {slide['number']}: {error}") from error
    with tempfile.TemporaryDirectory(prefix="render-", dir=folder) as scratch:
        scratch = Path(scratch)
        for index, (slide, timing) in enumerate(zip(slides, timeline, strict=True)):
            progress(index / len(slides), f"Rendering slide {index + 1} of {len(slides)}")
            segment = scratch / f"segment-{index:03}.mkv"
            run_command(
                [
                    ffmpeg(),
                    "-nostdin",
                    "-y",
                    "-v",
                    "error",
                    "-loop",
                    "1",
                    "-framerate",
                    str(FPS),
                    "-i",
                    str(folder / slide["image"]),
                    "-i",
                    str(folder / slide["audio"]),
                    "-vf",
                    f"scale={width}:{height}:force_original_aspect_ratio=decrease,pad={width}:{height}:(ow-iw)/2:(oh-ih)/2:color=white,setsar=1",
                    "-af",
                    "apad",
                    "-t",
                    str(timing["frames"] / FPS),
                    "-r",
                    str(FPS),
                    "-c:v",
                    "libx264",
                    "-tune",
                    "stillimage",
                    "-preset",
                    "veryfast",
                    "-crf",
                    "20",
                    "-pix_fmt",
                    "yuv420p",
                    "-c:a",
                    "pcm_s16le",
                    "-ar",
                    "48000",
                    "-threads",
                    "2",
                    str(segment),
                ],
                timeout=1800,
            )
        progress(0.96, "Joining slides and finalising audio")
        playlist = scratch / "segments.txt"
        playlist.write_text("\n".join(f"file 'segment-{index:03}.mkv'" for index in range(len(slides))))
        output = scratch / "video.mp4"
        run_command(
            [
                ffmpeg(),
                "-nostdin",
                "-y",
                "-v",
                "error",
                "-f",
                "concat",
                "-safe",
                "0",
                "-i",
                str(playlist),
                "-c:v",
                "copy",
                "-c:a",
                "aac",
                "-b:a",
                "192k",
                "-movflags",
                "+faststart",
                str(output),
            ],
            timeout=1800,
        )
        output.replace(folder / "video.mp4")
    return timeline


def demo_slides(folder):
    content = [
        (
            "Your slides.\nYour voice.",
            "A simpler way to share what you know.",
            "Welcome to SlideVoice. Turn your presentations into narrated videos, without manually editing a timeline.",
        ),
        (
            "Let the audio\nset the pace.",
            "Every slide gets exactly the time it needs.",
            "Each slide has its own audio clip. SlideVoice measures the recording and automatically matches the slide duration to it.",
        ),
        (
            "Create once.\nExplain beautifully.",
            "Review your script. Choose your voice. Export.",
            "This demo uses a Mac preview voice, not a clone of your voice. Connect your own voice when you are ready.",
        ),
    ]
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    heading = (
        ImageFont.truetype(font_path, 92) if Path(font_path).exists() else ImageFont.load_default(size=92)
    )
    small = ImageFont.truetype(font_path, 26) if Path(font_path).exists() else ImageFont.load_default(size=26)
    slides = []
    for position, (title, subtitle, script) in enumerate(content, 1):
        image = Image.new("RGB", (1600, 900), "#f5f3ed")
        drawing = ImageDraw.Draw(image)
        drawing.rounded_rectangle((1100, 130, 1460, 750), radius=150, fill="#dadfc9")
        for bar in range(13):
            bar_height = 50 + int(abs(math.sin(bar * 0.7)) * 160)
            horizontal = 1140 + bar * 22
            drawing.rounded_rectangle(
                (horizontal, 450 - bar_height / 2, horizontal + 10, 450 + bar_height / 2),
                radius=5,
                fill="#536647",
            )
        drawing.text((100, 90), "SLIDEVOICE / A LITTLE LESS BUSYWORK", fill="#536647", font=small)
        drawing.multiline_text((100, 260), title, fill="#252d24", font=heading, spacing=8)
        drawing.text((105, 580), subtitle, fill="#60675c", font=small)
        drawing.text((105, 790), f"{position:02} / 03", fill="#60675c", font=small)
        filename = f"slide-{position:03}.png"
        image.save(folder / filename)
        slide = slide_record(position, filename, title.replace("\n", " "), subtitle)
        slide["script"] = script
        slides.append(slide)
    return slides
