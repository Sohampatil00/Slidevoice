import base64
import json
import math
import shutil
import wave

import httpx
import pymupdf as fitz
import pytest
from PIL import Image
from pptx import Presentation

from app import core, media, providers
from app.core import WorkflowError


def endpoint(project, action):
    return f"/api/projects/{project['id']}/{action}"


@pytest.mark.parametrize("signal", ["silence", "quiet", "dc", "click"])
def test_inaudible_audio_is_rejected_without_replacing_recording(tmp_path, wav_bytes, signal):
    source = tmp_path / "source.wav"
    target = tmp_path / "existing.wav"
    target.write_bytes(wav_bytes)
    samples = {
        "silence": b"\x00\x00" * 48000,
        "quiet": b"\x01\x00\xff\xff" * 24000,
        "dc": b"\x00\x10" * 48000,
        "click": b"\xff\x7f" + b"\x00\x00" * 47999,
    }[signal]
    with wave.open(str(source), "wb") as recording:
        recording.setnchannels(1)
        recording.setsampwidth(2)
        recording.setframerate(48000)
        recording.writeframes(samples)
    with pytest.raises(media.SilentAudioError, match="silent or too quiet"):
        media.normalize_audio(source, target)
    assert target.read_bytes() == wav_bytes
    assert not list(tmp_path.glob("*-normalizing.wav"))


def test_export_blocks_preexisting_silent_recordings(client, project, wav_bytes, silent_wav_bytes):
    for slide in project["slides"]:
        response = client.post(
            endpoint(project, f"slides/{slide['id']}/audio"), files={"file": ("voice.wav", wav_bytes)}
        )
        assert response.status_code == 200
    saved = core.read_project(project["id"])
    (core.project_dir(project["id"]) / saved["slides"][1]["audio"]).write_bytes(silent_wav_bytes)
    client.post(endpoint(project, "render"), json={"resolution": "720p"})
    saved = core.read_project(project["id"])
    assert saved["job"]["status"] == "failed"
    assert "Slide 2" in saved["job"]["message"]
    assert "silent or too quiet" in saved["job"]["message"]
    assert not saved["export"]


def test_silent_cached_narration_is_regenerated(client, project, monkeypatch, wav_bytes, silent_wav_bytes):
    generated = []

    def synthesize(script, destination, *args, **kwargs):
        generated.append(script)
        destination.write_bytes(wav_bytes)
        return 1.0

    monkeypatch.setattr(providers, "synthesize", synthesize)
    body = {"provider": "preview"}
    client.post(endpoint(project, "audio"), json=body)
    saved = core.read_project(project["id"])
    assert saved["job"]["status"] == "complete"
    damaged = saved["slides"][1]
    (core.project_dir(project["id"]) / damaged["audio"]).write_bytes(silent_wav_bytes)
    generated.clear()
    client.post(endpoint(project, "audio"), json=body)
    saved = core.read_project(project["id"])
    assert saved["job"]["status"] == "complete"
    assert generated == [damaged["script"]]
    assert saved["slides"][1]["audio"] != damaged["audio"]


def test_demo_and_private_files(client, project):
    assert len(project["slides"]) == 3
    assert all(slide["script"] for slide in project["slides"])
    assert client.get(project["slides"][0]["image_url"]).status_code == 200
    assert client.get(endpoint(project, "files/project.json")).status_code == 404
    assert client.get(endpoint(project, "files/timings.json")).status_code == 404
    assert client.get("/api/projects/not-an-id").status_code == 400
    assert client.get("/api/projects").json()[0]["id"] == project["id"]


def test_loopback_and_csrf_protection(client):
    assert client.post("/api/demo", headers={"X-SlideVoice": ""}).status_code == 403
    assert client.post("/api/demo", headers={"Origin": "https://attacker.example"}).status_code == 403
    assert client.get("/", headers={"Host": "attacker.example"}).status_code == 403
    response = client.get("/")
    assert response.status_code == 200
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]


def test_health_does_not_expose_keys(client, monkeypatch):
    monkeypatch.setenv("INDIC_MIO_API_KEY", "secret-do-not-disclose")
    monkeypatch.setattr(providers, "indic_status", lambda: True)
    result = client.get("/api/health")
    assert result.json()["indic_mio"] is True
    assert "secret-do-not-disclose" not in result.text


def test_image_import_keeps_upload_order(client, png_bytes):
    response = client.post(
        "/api/projects",
        files=[
            ("files", ("02.png", png_bytes, "image/png")),
            ("files", ("01.png", png_bytes, "image/png")),
        ],
    )
    project = client.get(f"/api/projects/{response.json()['id']}").json()
    assert project["job"]["status"] == "complete"
    assert [slide["title"] for slide in project["slides"]] == ["02", "01"]


def test_pdf_import(client):
    with fitz.open() as document:
        page = document.new_page(width=960, height=540)
        page.insert_text((50, 50), "An important concept")
        payload = document.tobytes()
    response = client.post("/api/projects", files={"files": ("lesson.pdf", payload, "application/pdf")})
    project = client.get(f"/api/projects/{response.json()['id']}").json()
    assert project["job"]["status"] == "complete"
    assert project["slides"][0]["title"] == "An important concept"
    assert "important concept" in project["slides"][0]["text"]


def test_pptx_import(tmp_path):
    if not media.soffice():
        pytest.skip("LibreOffice unavailable")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Keep the right slide order"
    slide.placeholders[1].text = "The explanation follows the image."
    slide.notes_slide.notes_text_frame.text = "A useful example for narration."
    source = tmp_path / "lesson.pptx"
    presentation.save(source)
    result = media.import_document(source, tmp_path)
    assert len(result) == 1
    assert result[0]["title"] == "Keep the right slide order"
    assert "useful example" in result[0]["notes"]
    assert (tmp_path / result[0]["image"]).is_file()


@pytest.mark.parametrize("filename,payload", [("test.exe", b"binary"), ("empty.pdf", b"")])
def test_invalid_uploads(client, filename, payload):
    assert client.post("/api/projects", files={"files": (filename, payload)}).status_code == 400
    assert client.get("/api/projects").json() == []


def test_bad_pdf_failure_is_visible(client):
    response = client.post("/api/projects", files={"files": ("bad.pdf", b"not a pdf")})
    project = client.get(f"/api/projects/{response.json()['id']}").json()
    assert project["job"]["status"] == "failed"
    assert not core.project_lock(project["id"]).locked()


def test_reject_oversize_upload(client, png_bytes, monkeypatch):
    monkeypatch.setattr(core, "MAX_UPLOAD_BYTES", 10)
    assert client.post("/api/projects", files={"files": ("image.png", png_bytes)}).status_code == 400
    assert client.get("/api/projects").json() == []


def test_manual_audio_then_script_edit_invalidates(client, project, wav_bytes):
    audio_url = endpoint(project, "slides/slide-001/audio")
    response = client.post(audio_url, files={"file": ("recording.wav", wav_bytes)})
    assert response.status_code == 200
    result = response.json()
    slide = result["slides"][0]
    assert slide["duration"] == pytest.approx(1, abs=0.001)
    assert slide["audio_source"] == "uploaded"
    response = client.patch(
        endpoint(project, "slides/slide-001"),
        json={"title": "A better title", "script": "A new explanation."},
    )
    assert response.json()["slides"][0]["audio_ready"] is False
    assert client.get(slide["audio_url"]).status_code == 404


def test_invalid_audio_keeps_previous_recording(client, project, wav_bytes):
    response = client.post(
        endpoint(project, "slides/slide-001/audio"), files={"file": ("voice.wav", wav_bytes)}
    )
    saved = response.json()["slides"][0]["audio_url"]
    response = client.post(
        endpoint(project, "slides/slide-001/audio"), files={"file": ("bad.wav", b"bad audio")}
    )
    assert response.status_code == 400
    assert client.get(saved).status_code == 200


def test_project_lock_rejects_mutations(client, project):
    lock = core.project_lock(project["id"])
    lock.acquire()
    try:
        assert (
            client.patch(
                endpoint(project, "slides/slide-001"), json={"title": "Busy", "script": "Wait"}
            ).status_code
            == 409
        )
        assert client.post(endpoint(project, "render"), json={}).status_code == 409
        assert client.get(f"/api/projects/{project['id']}").status_code == 200
    finally:
        lock.release()


def test_voice_consent_and_configuration_required(client, project, monkeypatch):
    assert client.post(endpoint(project, "audio"), json={"provider": "indic_mio"}).status_code == 400
    assert (
        client.post(
            endpoint(project, "audio"),
            json={"provider": "indic_mio", "voice_id": "own-voice", "consent": True, "speed": 5},
        ).status_code
        == 422
    )
    assert client.post(endpoint(project, "scripts"), json={}).status_code == 400


def test_speech_cache_and_selective_regeneration(client, project, monkeypatch, wav_bytes):
    calls = []

    def fake_speech(script, destination, provider, voice_id, speed, reference_audio=None):
        calls.append(script)
        destination.write_bytes(wav_bytes)
        return 1.0

    monkeypatch.setattr(providers, "synthesize", fake_speech)
    body = {"provider": "indic_mio", "voice_id": "own-voice", "consent": True, "use_reference": False}
    client.post(endpoint(project, "audio"), json=body)
    assert len(calls) == 3
    client.post(endpoint(project, "audio"), json=body)
    assert len(calls) == 3
    client.patch(
        endpoint(project, "slides/slide-002"),
        json={"title": "Second slide", "script": "A changed explanation."},
    )
    client.post(endpoint(project, "audio"), json=body)
    assert len(calls) == 4
    assert calls[-1] == "A changed explanation."
    assert client.get(f"/api/projects/{project['id']}").json()["job"]["status"] == "complete"


def test_bulk_audio_preserves_uploaded_recording(client, project, monkeypatch, wav_bytes):
    client.patch(endpoint(project, "slides/slide-001"), json={"title": "Recorded manually", "script": ""})
    recorded = client.post(
        endpoint(project, "slides/slide-001/audio"), files={"file": ("recording.wav", wav_bytes)}
    ).json()["slides"][0]

    def fake_speech(script, destination, provider, voice_id, speed, reference_audio=None):
        destination.write_bytes(wav_bytes)
        return 1.0

    monkeypatch.setattr(providers, "synthesize", fake_speech)
    client.post(
        endpoint(project, "audio"),
        json={"provider": "indic_mio", "voice_id": "own", "consent": True, "use_reference": False},
    )
    result = client.get(f"/api/projects/{project['id']}").json()
    assert result["job"]["status"] == "complete"
    assert result["slides"][0]["audio_url"] == recorded["audio_url"]


def test_partial_provider_failure_keeps_completed_slides(client, project, monkeypatch, wav_bytes):
    calls = []

    def fake_speech(script, destination, provider, voice_id, speed, reference_audio=None):
        calls.append(script)
        if len(calls) == 2:
            raise WorkflowError("Indic-MIO: server disconnected")
        destination.write_bytes(wav_bytes)
        return 1.0

    monkeypatch.setattr(providers, "synthesize", fake_speech)
    client.post(
        endpoint(project, "audio"),
        json={"provider": "indic_mio", "voice_id": "own", "consent": True, "use_reference": False},
    )
    result = client.get(f"/api/projects/{project['id']}").json()
    assert result["job"]["status"] == "failed"
    assert "Slide 2:" in result["job"]["message"]
    assert result["slides"][0]["audio_ready"]
    assert not result["slides"][1]["audio_ready"]
    assert not core.project_lock(project["id"]).locked()


def test_script_generation_only_fills_blanks(client, project, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "mock-key")
    client.patch(endpoint(project, "slides/slide-002"), json={"title": "Second", "script": ""})
    calls = []

    def fake_script(slide, folder, titles, instructions):
        calls.append(slide["number"])
        return "A newly generated narration."

    monkeypatch.setattr(providers, "generate_script", fake_script)
    client.post(endpoint(project, "scripts"), json={})
    assert calls == [2]
    result = client.get(f"/api/projects/{project['id']}").json()
    assert result["slides"][0]["script"] == project["slides"][0]["script"]


def test_timeline_rounds_up_without_truncating_audio():
    slides = [
        {"number": 1, "title": "One", "script": "A", "audio": "one.wav", "duration": 1.011},
        {"number": 2, "title": "Two", "script": "B", "audio": "two.wav", "duration": 2.002},
    ]
    timeline = media.timeline_for(slides, 0.5)
    assert timeline[0]["frames"] == math.ceil(1.511 * 30)
    assert timeline[0]["end"] >= 1.511
    assert timeline[1]["start"] == timeline[0]["end"]
    assert timeline[1]["end"] == pytest.approx((math.ceil(1.511 * 30) + math.ceil(2.502 * 30)) / 30)
    with pytest.raises(WorkflowError):
        media.timeline_for(slides, float("nan"))


def test_real_mp4_export_and_timings(client, project, wav_bytes, tmp_path):
    for slide in project["slides"]:
        response = client.post(
            endpoint(project, f"slides/{slide['id']}/audio"), files={"file": ("voice.wav", wav_bytes)}
        )
        assert response.status_code == 200
    response = client.post(endpoint(project, "render"), json={"pause": 0.3, "resolution": "720p"})
    assert response.status_code == 200
    result = client.get(f"/api/projects/{project['id']}").json()
    assert result["job"]["status"] == "complete", result["job"]["message"]
    assert result["export"]["duration"] == pytest.approx(3.9)
    video_response = client.get(result["export"]["video_url"])
    assert video_response.status_code == 200
    assert len(video_response.content) > 1000
    timings = client.get(result["export"]["timings_url"]).json()["slides"]
    assert timings[1]["start"] == pytest.approx(1.3)
    assert timings[2]["end"] == pytest.approx(3.9)
    folder = core.project_dir(project["id"])
    decoded = tmp_path / "decoded.wav"
    media.normalize_audio(folder / "video.mp4", decoded)
    with wave.open(str(decoded), "rb") as recording:
        assert recording.getnframes() / recording.getframerate() == pytest.approx(3.9, abs=0.08)
    screenshot = tmp_path / "frame.png"
    media.run_command(
        [
            media.ffmpeg(),
            "-nostdin",
            "-y",
            "-v",
            "error",
            "-i",
            str(folder / "video.mp4"),
            "-frames:v",
            "1",
            str(screenshot),
        ]
    )
    with Image.open(screenshot) as frame:
        assert frame.size == (1280, 720)
    client.patch(endpoint(project, "slides/slide-001"), json={"title": "Edited", "script": "Changed"})
    assert client.get(result["export"]["video_url"]).status_code == 404


def test_export_requires_all_audio(client, project):
    client.post(endpoint(project, "render"), json={})
    result = client.get(f"/api/projects/{project['id']}").json()
    assert result["job"]["status"] == "failed"
    assert "slide 1" in result["job"]["message"]
    assert result["export"] is None


def test_restart_marks_interrupted_job_failed(client, project):
    from fastapi.testclient import TestClient

    from app.main import app

    current = core.read_project(project["id"])
    current["job"]["status"] = "running"
    core.save_project(current)
    with TestClient(app, base_url="http://127.0.0.1") as restarted:
        result = restarted.get(f"/api/projects/{project['id']}").json()
    assert result["job"]["status"] == "failed"
    assert "restarted" in result["job"]["message"]


def test_provider_error_does_not_leak_response_secrets():
    with pytest.raises(WorkflowError) as caught:
        providers.provider_error(httpx.Response(401, text="secret-api-key"), "Indic-MIO")
    assert "secret-api-key" not in str(caught.value)


def test_claude_request_includes_slide_visual_and_notes(monkeypatch, project):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "mock-key")
    requests = []
    real_client = httpx.Client

    def respond(request):
        requests.append(json.loads(request.content))
        return httpx.Response(
            200,
            json={"stop_reason": "end_turn", "content": [{"type": "text", "text": "A clear explanation."}]},
        )

    monkeypatch.setattr(
        providers.httpx, "Client", lambda **kwargs: real_client(transport=httpx.MockTransport(respond))
    )
    slide = project["slides"][0]
    result = providers.generate_script(slide, core.project_dir(project["id"]), ["An outline"], "Be concise")
    assert result == "A clear explanation."
    assert requests[0]["messages"][0]["content"][0]["source"]["media_type"] == "image/jpeg"
    assert "speaker_notes" in requests[0]["messages"][0]["content"][1]["text"]


def test_indic_payload_and_audio_decoding(monkeypatch, project, wav_bytes):
    monkeypatch.setenv("INDIC_MIO_API_KEY", "mock-key")
    requests = []
    real_client = httpx.Client

    def respond(request):
        requests.append(request)
        return httpx.Response(200, json={"audio": base64.b64encode(wav_bytes).decode(), "token_count": 25})

    monkeypatch.setattr(
        providers.httpx, "Client", lambda **kwargs: real_client(transport=httpx.MockTransport(respond))
    )
    target = core.project_dir(project["id"]) / "test.wav"
    duration = providers.synthesize("A test explanation.", target, "indic_mio", "personal-voice", 1.0)
    assert duration == pytest.approx(1)
    payload = json.loads(requests[0].content)
    assert payload["reference"] == {"type": "preset", "preset_id": "personal-voice"}
    assert payload["llm"]["model"] == "SPRINGLab/Indic-Mio"
    assert requests[0].url.path == "/v1/tts"
    assert requests[0].headers["authorization"] == "Bearer mock-key"
    assert target.is_file()


def test_local_indian_english_preview(project, monkeypatch):
    if not shutil.which("say"):
        pytest.skip("macOS speech unavailable")
    monkeypatch.setenv("MACOS_PREVIEW_VOICE", "Aman")
    target = core.project_dir(project["id"]) / "preview.wav"
    duration = providers.synthesize("Welcome to your presentation.", target, "preview", "", 1.0)
    assert duration > 0.5
